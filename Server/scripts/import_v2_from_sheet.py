#!/usr/bin/env python
"""
Import data from employee/asset spreadsheet into AMS schema.

Order of operations per row:
1) employees
2) assets
3) asset_components
4) asset_assignments

IMPORTANT BUSINESS RULE:
ERP status is mapped to employees.is_active only.
ERP status is never mapped to assets.status.
Asset status is derived from open/closed assignments in DB triggers.

Operational rule:
This importer is for bootstrap/refresh workflows with service role.
Day-to-day assignment and return actions must go through RPCs:
fn_assign_asset / fn_return_asset.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import uuid
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from supabase import Client, create_client

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.settings import settings  # noqa: E402


ALIASES: dict[str, list[str]] = {
    "employee_code": ["employee_id", "employee_code", "emp_id", "employee"],
    "employee_name": ["employee_name", "employee", "name", "employee_full_name"],
    "employee_email": ["email", "employee_email", "mail", "official_email"],
    "department": ["department", "dept", "function"],
    "erp_status": ["erp_status", "erp status", "status_erp", "status"],
    "has_google_profile": ["has_google_profile", "google_profile", "has_google"],
    "has_erp_access": ["has_erp_access", "erp_access", "erp_login_access"],
    "asset_id": ["asset_id", "asset id", "asset_tag", "asset code"],
    "system_type": ["system_type", "system type", "asset_type", "type", "category"],
    "manufacturer": ["manufacturer", "brand", "make", "oem"],
    "model": ["model", "model_name"],
    "serial_number": ["serial_number", "serial no", "serial", "serial_no", "sn"],
    "location": ["location", "location_code", "site", "branch"],
    "assigned_to": ["assigned_to", "assigned to", "custodian", "current_custodian"],
    "assigned_at": ["assigned_at", "assigned_date", "issued_date", "issue_date"],
    "sim_previously_used_by": [
        "sim_previously_used_by",
        "sim previously used by",
        "previous_sim_user",
        "previously_used_by",
    ],
}

CUSTOM_FIELD_KEY_REMAP: dict[str, str] = {
    "processor_name": "processor",
    "cpu": "processor",
    "gen": "generation",
    "ram": "ram_gb",
    "ramgb": "ram_gb",
    "wifi_mac": "mac_wifi",
    "mac_address_wifi": "mac_wifi",
    "lan_mac": "mac_lan",
    "mac_address_lan": "mac_lan",
    "storage": "storage_gb",
    "storage_capacity_gb": "storage_gb",
}

CATEGORY_SYNONYMS: dict[str, list[str]] = {
    "laptop": ["laptop", "notebook"],
    "desktop": ["desktop", "workstation", "pc"],
    "sim": ["sim", "sim card", "data sim"],
    "pen-drive": ["pen drive", "pendrive", "pen-drive", "usb drive", "usb"],
    "monitor": ["monitor", "display", "lcd", "led"],
    "networking": ["networking", "router", "switch", "access point", "firewall"],
}

COMPONENT_FIELDS: list[dict[str, Any]] = [
    {
        "component_type": "ssd",
        "serial_aliases": ["ssd_serial", "ssd_serial_number", "ssd_sn"],
        "meta_aliases": ["ssd_capacity_gb", "ssd_size_gb", "ssd_model"],
    },
    {
        "component_type": "hdd",
        "serial_aliases": ["hdd_serial", "hdd_serial_number", "hdd_sn"],
        "meta_aliases": ["hdd_capacity_gb", "hdd_size_gb", "hdd_model"],
    },
    {
        "component_type": "screen",
        "serial_aliases": ["screen_serial", "screen_serial_number", "display_serial"],
        "meta_aliases": ["screen_size_inch", "resolution", "panel_type"],
    },
]


@dataclass
class ImportSummary:
    rows_seen: int = 0
    rows_with_assets: int = 0
    rows_skipped: int = 0
    employees_upserted: int = 0
    assets_upserted: int = 0
    components_upserted: int = 0
    assignments_inserted: int = 0
    history_assignments_inserted: int = 0


class ImportErrorWithContext(Exception):
    pass


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "" or text.lower() in {"na", "n/a", "none", "null", "nan"}:
        return None
    return text


def parse_bool(value: Any) -> bool | None:
    text = clean_text(value)
    if text is None:
        return None
    normalized = text.lower()
    if normalized in {"1", "true", "yes", "y", "active", "enabled"}:
        return True
    if normalized in {"0", "false", "no", "n", "inactive", "disabled"}:
        return False
    return None


def parse_is_active_from_erp(value: Any) -> bool:
    text = clean_text(value)
    if text is None:
        return True
    normalized = text.lower()
    inactive_tokens = [
        "inactive",
        "terminated",
        "disabled",
        "left",
        "exit",
        "resigned",
        "deactivated",
        "not active",
    ]
    return not any(token in normalized for token in inactive_tokens)


def parse_datetime(value: Any) -> str:
    text = clean_text(value)
    if text is None:
        return datetime.now(timezone.utc).isoformat()

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(tzinfo=timezone.utc).isoformat()
        except ValueError:
            continue

    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc).isoformat()
    except ValueError:
        raise ValueError(f"Invalid datetime value: {text}")


def normalize_slug(value: str) -> str:
    cleaned = normalize_header(value).replace("_", "-")
    if not cleaned:
        return "uncategorized"

    for slug, synonyms in CATEGORY_SYNONYMS.items():
        if cleaned == slug or cleaned in [normalize_header(s).replace("_", "-") for s in synonyms]:
            return slug
    return cleaned


def parse_location(value: str | None) -> tuple[str | None, str | None]:
    text = clean_text(value)
    if text is None:
        return None, None

    match = re.search(r"\(([^)]+)\)$", text)
    if match:
        code = normalize_location_code(match.group(1))
        name = text[: match.start()].strip() or text
        return code, name

    maybe_code = normalize_location_code(text)
    if re.fullmatch(r"[A-Z0-9]+(?:-[A-Z0-9]+)+", maybe_code or ""):
        return maybe_code, text

    code = normalize_location_code(text)
    return code, text


def normalize_location_code(value: str | None) -> str | None:
    text = clean_text(value)
    if text is None:
        return None
    code = re.sub(r"[^A-Za-z0-9]+", "-", text.upper()).strip("-")
    return code or None


def deterministic_code(prefix: str, seed: str, size: int = 8) -> str:
    digest = hashlib.sha1(seed.encode("utf-8")).hexdigest()[:size].upper()
    return f"{prefix}-{digest}"


class SpreadsheetLoader:
    def __init__(self, path: Path, sheet_name: str | None = None):
        self.path = path
        self.sheet_name = sheet_name

    def load(self) -> list[dict[str, str | None]]:
        suffix = self.path.suffix.lower()
        if suffix == ".csv":
            return self._load_csv()
        if suffix in {".xlsx", ".xlsm"}:
            return self._load_xlsx()
        raise ImportErrorWithContext(f"Unsupported file format: {self.path.name}")

    def _load_csv(self) -> list[dict[str, str | None]]:
        rows: list[dict[str, str | None]] = []
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                normalized = {normalize_header(k): clean_text(v) for k, v in row.items() if k}
                rows.append(normalized)
        return rows

    def _load_xlsx(self) -> list[dict[str, str | None]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportErrorWithContext(
                "openpyxl is required for .xlsx import. Install with: pip install openpyxl"
            ) from exc

        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb[self.sheet_name] if self.sheet_name and self.sheet_name in wb.sheetnames else wb.active

        iterator = ws.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            return []

        normalized_headers = [normalize_header(h) for h in headers]
        rows: list[dict[str, str | None]] = []
        for values in iterator:
            row: dict[str, str | None] = {}
            for idx, header in enumerate(normalized_headers):
                if not header:
                    continue
                value = values[idx] if idx < len(values) else None
                row[header] = clean_text(value)
            rows.append(row)
        return rows


class V2Importer:
    def __init__(self, client: Client, validate_only: bool = False):
        self.client = client
        self.validate_only = validate_only
        self.summary = ImportSummary()
        self.warnings: list[str] = []
        self.errors: list[str] = []

        self.department_cache: dict[str, str] = {}
        self.manufacturer_cache: dict[str, str] = {}
        self.location_cache: dict[str, str] = {}
        self.category_cache: dict[str, str] = {}
        self.custom_fields_by_category: dict[str, set[str]] = {}

        self.import_employee_codes: set[str] = set()
        self.import_asset_keys: set[str] = set()

    def run(self, rows: list[dict[str, str | None]]) -> dict[str, Any]:
        self._warm_reference_cache()

        for idx, row in enumerate(rows, start=2):
            self.summary.rows_seen += 1
            try:
                self._process_row(row, idx)
            except Exception as exc:  # noqa: BLE001
                self.errors.append(f"Row {idx}: {exc}")

        checks = self._post_import_checks()

        return {
            "summary": self.summary.__dict__,
            "import_distinct": {
                "employees_in_sheet": len(self.import_employee_codes),
                "assets_in_sheet": len(self.import_asset_keys),
            },
            "checks": checks,
            "warnings": self.warnings,
            "errors": self.errors,
            "validate_only": self.validate_only,
        }

    def _warm_reference_cache(self) -> None:
        for row in self.client.table("asset_categories").select("id,slug").execute().data or []:
            self.category_cache[row["slug"]] = row["id"]

        category_slug_by_id: dict[str, str] = {v: k for k, v in self.category_cache.items()}
        field_rows = (
            self.client.table("custom_field_definitions")
            .select("category_id,field_key")
            .execute()
            .data
            or []
        )
        for row in field_rows:
            slug = category_slug_by_id.get(row["category_id"])
            if not slug:
                continue
            self.custom_fields_by_category.setdefault(slug, set()).add(row["field_key"])

    def _process_row(self, row: dict[str, str | None], row_number: int) -> None:
        employee_id, employee_name = self._upsert_employee(row)
        asset_id = self._upsert_asset(row, row_number)

        if not asset_id:
            self.summary.rows_skipped += 1
            return

        self.summary.rows_with_assets += 1

        manufacturer_id = self._get_manufacturer_id(self._value(row, "manufacturer"))
        self._upsert_components(row, asset_id, manufacturer_id)
        self._upsert_assignments(row, asset_id, employee_id, employee_name)

    def _value(self, row: dict[str, str | None], alias_group: str) -> str | None:
        for alias in ALIASES[alias_group]:
            key = normalize_header(alias)
            value = clean_text(row.get(key))
            if value is not None:
                return value
        return None

    def _upsert_employee(self, row: dict[str, str | None]) -> tuple[str | None, str | None]:
        employee_code = self._value(row, "employee_code")
        employee_name = self._value(row, "employee_name") or self._value(row, "assigned_to")
        email = self._value(row, "employee_email")

        if not employee_code and not employee_name and not email:
            return None, None

        if employee_code is None:
            seed = email or employee_name or str(uuid.uuid4())
            employee_code = deterministic_code("AUTO", seed)

        department_name = self._value(row, "department")
        department_id = self._get_department_id(department_name)

        erp_status = self._value(row, "erp_status")
        is_active = parse_is_active_from_erp(erp_status)

        metadata: dict[str, Any] = {}
        has_google_profile = parse_bool(self._value(row, "has_google_profile"))
        has_erp_access = parse_bool(self._value(row, "has_erp_access"))
        if has_google_profile is not None:
            metadata["has_google_profile"] = has_google_profile
        if has_erp_access is not None:
            metadata["has_erp_access"] = has_erp_access

        if employee_code:
            self.import_employee_codes.add(employee_code)

        payload = {
            "employee_code": employee_code,
            "name": employee_name or employee_code,
            "email": email,
            "department_id": department_id,
            "is_active": is_active,
            "metadata": metadata,
        }

        if self.validate_only:
            self.summary.employees_upserted += 1
            return deterministic_code("EMP", employee_code), payload["name"]

        self.client.table("employees").upsert(payload, on_conflict="employee_code").execute()
        row_data = (
            self.client.table("employees")
            .select("id,name")
            .eq("employee_code", employee_code)
            .limit(1)
            .execute()
            .data
        )
        if not row_data:
            raise ImportErrorWithContext(f"Unable to fetch employee after upsert: {employee_code}")

        self.summary.employees_upserted += 1
        return row_data[0]["id"], row_data[0]["name"]

    def _upsert_asset(self, row: dict[str, str | None], row_number: int) -> str | None:
        system_type = self._value(row, "system_type")
        if system_type is None:
            self.warnings.append(f"Row {row_number}: missing system type, asset skipped")
            return None

        category_slug = normalize_slug(system_type)
        category_id = self._get_category_id(category_slug, system_type)

        manufacturer_name = self._value(row, "manufacturer")
        manufacturer_id = self._get_manufacturer_id(manufacturer_name)

        location_code, location_name = parse_location(self._value(row, "location"))
        location_id = self._get_location_id(location_code, location_name)

        asset_tag = self._value(row, "asset_id")
        serial_number = self._value(row, "serial_number")

        if asset_tag:
            self.import_asset_keys.add(asset_tag)
        elif serial_number:
            self.import_asset_keys.add(serial_number)
        else:
            self.warnings.append(
                f"Row {row_number}: missing both Asset_Id and serial_number; cannot upsert idempotently"
            )
            return None

        custom_fields = self._extract_custom_fields(row, category_slug)

        payload = {
            "asset_tag": asset_tag,
            "category_id": category_id,
            "manufacturer_id": manufacturer_id,
            "model": self._value(row, "model"),
            "serial_number": serial_number,
            "location_id": location_id,
            "custom_fields": custom_fields,
            "metadata": {"import_source": "v2_sheet"},
        }

        conflict_key = "asset_tag" if asset_tag else "serial_number"

        if self.validate_only:
            self.summary.assets_upserted += 1
            return deterministic_code("AST", asset_tag or serial_number or str(row_number))

        self.client.table("assets").upsert(payload, on_conflict=conflict_key).execute()

        query = self.client.table("assets").select("id")
        if asset_tag:
            query = query.eq("asset_tag", asset_tag)
        else:
            query = query.eq("serial_number", serial_number)

        row_data = query.limit(1).execute().data
        if not row_data:
            raise ImportErrorWithContext(f"Unable to fetch asset after upsert at row {row_number}")

        self.summary.assets_upserted += 1
        return row_data[0]["id"]

    def _upsert_components(self, row: dict[str, str | None], asset_id: str, fallback_manufacturer_id: str | None) -> None:
        for component_spec in COMPONENT_FIELDS:
            serial = self._first_by_keys(row, component_spec["serial_aliases"])
            if serial is None:
                continue

            metadata: dict[str, Any] = {}
            for meta_key in component_spec["meta_aliases"]:
                normalized = normalize_header(meta_key)
                value = clean_text(row.get(normalized))
                if value is not None:
                    metadata[normalized] = value

            payload = {
                "asset_id": asset_id,
                "component_type": component_spec["component_type"],
                "manufacturer_id": fallback_manufacturer_id,
                "serial_number": serial,
                "metadata": metadata,
            }

            if self.validate_only:
                self.summary.components_upserted += 1
                continue

            self.client.table("asset_components").upsert(payload, on_conflict="serial_number").execute()
            self.summary.components_upserted += 1

    def _upsert_assignments(
        self,
        row: dict[str, str | None],
        asset_id: str,
        employee_id: str | None,
        employee_name: str | None,
    ) -> None:
        assigned_at = parse_datetime(self._value(row, "assigned_at"))

        if employee_id is not None:
            self._ensure_current_assignment(asset_id, employee_id, assigned_at)

        history_raw = self._value(row, "sim_previously_used_by")
        if history_raw is None:
            return

        history_names = [clean_text(part) for part in re.split(r"[,;|]", history_raw)]
        filtered_names = [name for name in history_names if name and name != employee_name]

        for idx, name in enumerate(filtered_names):
            hist_code = deterministic_code("HIST", name)
            hist_payload = {
                "employee_code": hist_code,
                "name": name,
                "is_active": False,
                "metadata": {"source": "sim_history"},
            }

            if self.validate_only:
                hist_employee_id = deterministic_code("EMP", hist_code)
            else:
                self.client.table("employees").upsert(hist_payload, on_conflict="employee_code").execute()
                emp_data = (
                    self.client.table("employees")
                    .select("id")
                    .eq("employee_code", hist_code)
                    .limit(1)
                    .execute()
                    .data
                )
                if not emp_data:
                    continue
                hist_employee_id = emp_data[0]["id"]

            start_dt = datetime.fromisoformat(assigned_at.replace("Z", "+00:00")) - timedelta(days=(idx + 2) * 30)
            end_dt = start_dt + timedelta(days=20)
            self._ensure_history_assignment(asset_id, hist_employee_id, start_dt.isoformat(), end_dt.isoformat())

    def _ensure_current_assignment(self, asset_id: str, employee_id: str, assigned_at: str) -> None:
        if self.validate_only:
            self.summary.assignments_inserted += 1
            return

        open_rows = (
            self.client.table("asset_assignments")
            .select("id,employee_id")
            .eq("asset_id", asset_id)
            .is_("returned_at", "null")
            .execute()
            .data
            or []
        )

        if len(open_rows) == 1 and open_rows[0]["employee_id"] == employee_id:
            return

        for open_row in open_rows:
            self.client.table("asset_assignments").update({"returned_at": assigned_at}).eq(
                "id", open_row["id"]
            ).execute()

        self.client.table("asset_assignments").insert(
            {
                "asset_id": asset_id,
                "employee_id": employee_id,
                "assigned_at": assigned_at,
                "returned_at": None,
                "source": "import_current",
                "notes": "Imported current custodian",
            }
        ).execute()
        self.summary.assignments_inserted += 1

    def _ensure_history_assignment(
        self,
        asset_id: str,
        employee_id: str,
        assigned_at: str,
        returned_at: str,
    ) -> None:
        if self.validate_only:
            self.summary.history_assignments_inserted += 1
            return

        existing = (
            self.client.table("asset_assignments")
            .select("id")
            .eq("asset_id", asset_id)
            .eq("employee_id", employee_id)
            .eq("source", "sim_history")
            .limit(1)
            .execute()
            .data
        )
        if existing:
            return

        self.client.table("asset_assignments").insert(
            {
                "asset_id": asset_id,
                "employee_id": employee_id,
                "assigned_at": assigned_at,
                "returned_at": returned_at,
                "source": "sim_history",
                "notes": "Imported from SIM previously used by",
            }
        ).execute()
        self.summary.history_assignments_inserted += 1

    def _extract_custom_fields(self, row: dict[str, str | None], category_slug: str) -> dict[str, Any]:
        allowed_keys = self.custom_fields_by_category.get(category_slug, set())
        if not allowed_keys:
            return {}

        custom_fields: dict[str, Any] = {}
        for raw_key, raw_value in row.items():
            value = clean_text(raw_value)
            if value is None:
                continue

            normalized_key = normalize_header(raw_key)
            normalized_key = CUSTOM_FIELD_KEY_REMAP.get(normalized_key, normalized_key)

            if normalized_key in allowed_keys:
                custom_fields[normalized_key] = self._coerce_custom_value(value)

        return custom_fields

    @staticmethod
    def _coerce_custom_value(value: str) -> Any:
        bool_value = parse_bool(value)
        if bool_value is not None:
            return bool_value

        if re.fullmatch(r"\d+", value):
            return int(value)
        if re.fullmatch(r"\d+\.\d+", value):
            return float(value)
        return value

    def _first_by_keys(self, row: dict[str, str | None], keys: list[str]) -> str | None:
        for key in keys:
            value = clean_text(row.get(normalize_header(key)))
            if value is not None:
                return value
        return None

    def _get_department_id(self, department_name: str | None) -> str | None:
        name = clean_text(department_name)
        if name is None:
            return None

        if name in self.department_cache:
            return self.department_cache[name]

        if self.validate_only:
            fake_id = deterministic_code("DEPT", name)
            self.department_cache[name] = fake_id
            return fake_id

        self.client.table("departments").upsert({"name": name}, on_conflict="name").execute()
        rows = self.client.table("departments").select("id").eq("name", name).limit(1).execute().data
        if not rows:
            return None
        dept_id = rows[0]["id"]
        self.department_cache[name] = dept_id
        return dept_id

    def _get_manufacturer_id(self, manufacturer_name: str | None) -> str | None:
        name = clean_text(manufacturer_name)
        if name is None:
            return None

        if name in self.manufacturer_cache:
            return self.manufacturer_cache[name]

        if self.validate_only:
            fake_id = deterministic_code("MFG", name)
            self.manufacturer_cache[name] = fake_id
            return fake_id

        self.client.table("manufacturers").upsert({"name": name}, on_conflict="name").execute()
        rows = (
            self.client.table("manufacturers").select("id").eq("name", name).limit(1).execute().data
        )
        if not rows:
            return None
        mfg_id = rows[0]["id"]
        self.manufacturer_cache[name] = mfg_id
        return mfg_id

    def _get_location_id(self, location_code: str | None, location_name: str | None) -> str | None:
        code = normalize_location_code(location_code or location_name)
        name = clean_text(location_name) or code
        if code is None:
            return None

        if code in self.location_cache:
            return self.location_cache[code]

        payload = {"code": code, "name": name}
        if self.validate_only:
            fake_id = deterministic_code("LOC", code)
            self.location_cache[code] = fake_id
            return fake_id

        self.client.table("locations").upsert(payload, on_conflict="code").execute()
        rows = self.client.table("locations").select("id").eq("code", code).limit(1).execute().data
        if not rows:
            return None
        loc_id = rows[0]["id"]
        self.location_cache[code] = loc_id
        return loc_id

    def _get_category_id(self, slug: str, label_seed: str) -> str:
        if slug in self.category_cache:
            return self.category_cache[slug]

        name = clean_text(label_seed) or slug.replace("-", " ").title()
        payload = {"slug": slug, "name": name}

        if self.validate_only:
            fake_id = deterministic_code("CAT", slug)
            self.category_cache[slug] = fake_id
            return fake_id

        self.client.table("asset_categories").upsert(payload, on_conflict="slug").execute()
        rows = (
            self.client.table("asset_categories")
            .select("id")
            .eq("slug", slug)
            .limit(1)
            .execute()
            .data
        )
        if not rows:
            raise ImportErrorWithContext(f"Failed to resolve category id for slug={slug}")

        category_id = rows[0]["id"]
        self.category_cache[slug] = category_id
        return category_id

    def _post_import_checks(self) -> dict[str, Any]:
        checks: dict[str, Any] = {
            "sheet_vs_db": {},
            "assignment_integrity": {},
            "inactive_employee_open_assignments": [],
        }

        if self.validate_only:
            checks["note"] = "Validation-only mode skips DB integrity queries"
            return checks

        employees_count = self.client.table("employees").select("id", count="exact").execute().count
        assets_count = self.client.table("assets").select("id", count="exact").execute().count

        checks["sheet_vs_db"] = {
            "employees_in_sheet_distinct": len(self.import_employee_codes),
            "assets_in_sheet_distinct": len(self.import_asset_keys),
            "employees_total_in_db": employees_count,
            "assets_total_in_db": assets_count,
        }

        inventory_rows = (
            self.client.table("v_asset_inventory")
            .select(
                "id,asset_tag,status,assignment_id,current_employee_code,current_employee_is_active,current_employee_erp_active"
            )
            .execute()
            .data
            or []
        )

        assigned_without_open: list[str] = []
        open_but_not_assigned: list[str] = []
        inactive_holders: list[dict[str, Any]] = []

        for row in inventory_rows:
            asset_label = row.get("asset_tag") or row.get("id")
            status = row.get("status")
            has_open_assignment = row.get("assignment_id") is not None

            if status == "assigned" and not has_open_assignment:
                assigned_without_open.append(asset_label)
            if status != "assigned" and has_open_assignment:
                open_but_not_assigned.append(asset_label)
            if has_open_assignment and row.get("current_employee_erp_active") is False:
                inactive_holders.append(
                    {
                        "asset": asset_label,
                        "employee_code": row.get("current_employee_code"),
                    }
                )

        anomalies = self.client.table("v_assignment_anomalies").select("asset_id,asset_tag").execute().data or []

        checks["assignment_integrity"] = {
            "db_anomaly_rows": len(anomalies),
            "assets_assigned_without_open_assignment": assigned_without_open,
            "assets_open_assignment_but_status_not_assigned": open_but_not_assigned,
        }
        checks["inactive_employee_open_assignments"] = inactive_holders
        return checks


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import AMS data from CSV/XLSX")
    parser.add_argument("--file", required=True, help="Path to spreadsheet (.csv or .xlsx)")
    parser.add_argument("--sheet", help="Worksheet name for .xlsx files")
    parser.add_argument("--url", default=settings.SUPABASE_URL, help="Supabase URL")
    parser.add_argument("--key", default=settings.SUPABASE_KEY, help="Supabase service key")
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="Parse and map data without writing rows",
    )
    parser.add_argument(
        "--report-file",
        help="Optional path to save JSON report",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if not args.url or not args.key:
        raise ImportErrorWithContext("Supabase URL/key is required. Set .env or pass --url and --key")

    file_path = Path(args.file)
    if not file_path.exists():
        raise ImportErrorWithContext(f"Input file not found: {file_path}")

    rows = SpreadsheetLoader(file_path, args.sheet).load()
    client = create_client(args.url, args.key)

    importer = V2Importer(client=client, validate_only=args.validate_only)
    report = importer.run(rows)

    report_json = json.dumps(report, indent=2)
    print(report_json)

    if args.report_file:
        Path(args.report_file).write_text(report_json, encoding="utf-8")


if __name__ == "__main__":
    main()
